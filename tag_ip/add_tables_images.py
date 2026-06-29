#!/usr/bin/env python3
"""Insert actual images into ftc_corrige.docx before each figure caption.

Also creates real Word tables and fixes formatting issues.
Processes in reverse order (bottom to top) so insertions don't shift positions.
"""

import os
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image as PILImage

SRC = "/Users/fitahiana/Desktop/ftc_corrige.docx"
DST = "/Users/fitahiana/Desktop/ftc_corrige.docx"
IMG_DIR = "/Users/fitahiana/Desktop/fitahiana/"

FONT_NAME = "Times New Roman"


def make_run(text, bold=False, italic=False, size_pt=12, color=None, font_name=FONT_NAME):
    r = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rPr.append(rFonts)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), str(int(size_pt * 2)))
    rPr.append(sz)
    if bold:
        rPr.append(OxmlElement('w:b'))
    if italic:
        rPr.append(OxmlElement('w:i'))
    if color:
        c = OxmlElement('w:color')
        c.set(qn('w:val'), color)
        rPr.append(c)
    r.append(rPr)
    t = OxmlElement('w:t')
    t.text = text
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    r.append(t)
    return r


def make_empty_para():
    p = OxmlElement('w:p')
    p.append(OxmlElement('w:pPr'))
    r = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), FONT_NAME)
    rFonts.set(qn('w:hAnsi'), FONT_NAME)
    rPr.append(rFonts)
    r.append(rPr)
    t = OxmlElement('w:t')
    t.text = ""
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    r.append(t)
    p.append(r)
    return p


def build_image_para(doc, img_path, width_inches=5.0):
    rId, image = doc.part.get_or_add_image(img_path)

    with PILImage.open(img_path) as pil:
        img_width_px, img_height_px = pil.size

    width_emu = int(width_inches * 914400)
    aspect = img_height_px / img_width_px
    height_emu = int(width_emu * aspect)

    ns_wp = 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'

    p = OxmlElement('w:p')
    pPr = OxmlElement('w:pPr')
    jc = OxmlElement('w:jc')
    jc.set(qn('w:val'), 'center')
    pPr.append(jc)
    p.append(pPr)

    r = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), FONT_NAME)
    rFonts.set(qn('w:hAnsi'), FONT_NAME)
    rPr.append(rFonts)
    r.append(rPr)

    drawing = OxmlElement('w:drawing')
    inline = OxmlElement('wp:inline')
    inline.set(f'{{{ns_wp}}}distT', '0')
    inline.set(f'{{{ns_wp}}}distB', '0')
    inline.set(f'{{{ns_wp}}}distL', '0')
    inline.set(f'{{{ns_wp}}}distR', '0')

    extent = OxmlElement('wp:extent')
    extent.set('cx', str(width_emu))
    extent.set('cy', str(height_emu))
    inline.append(extent)

    effectExtent = OxmlElement('wp:effectExtent')
    effectExtent.set('l', '0')
    effectExtent.set('t', '0')
    effectExtent.set('r', '0')
    effectExtent.set('b', '0')
    inline.append(effectExtent)

    docPr = OxmlElement('wp:docPr')
    docPr.set('id', str(abs(hash(img_path)) % 100000))
    docPr.set('name', os.path.basename(img_path))
    inline.append(docPr)

    graphic = OxmlElement('a:graphic')
    graphicData = OxmlElement('a:graphicData')
    graphicData.set('uri', 'http://schemas.openxmlformats.org/drawingml/2006/picture')

    pic = OxmlElement('pic:pic')
    nvPicPr = OxmlElement('pic:nvPicPr')
    cNvPr = OxmlElement('pic:cNvPr')
    cNvPr.set('id', '0')
    cNvPr.set('name', os.path.basename(img_path))
    nvPicPr.append(cNvPr)
    nvPicPr.append(OxmlElement('pic:cNvPicPr'))
    pic.append(nvPicPr)

    blipFill = OxmlElement('pic:blipFill')
    blip = OxmlElement('a:blip')
    blip.set(qn('r:embed'), rId)
    blipFill.append(blip)
    stretch = OxmlElement('a:stretch')
    stretch.append(OxmlElement('a:fillRect'))
    blipFill.append(stretch)
    pic.append(blipFill)

    spPr = OxmlElement('pic:spPr')
    xfrm = OxmlElement('a:xfrm')
    off = OxmlElement('a:off')
    off.set('x', '0')
    off.set('y', '0')
    xfrm.append(off)
    ext = OxmlElement('a:ext')
    ext.set('cx', str(width_emu))
    ext.set('cy', str(height_emu))
    xfrm.append(ext)
    spPr.append(xfrm)
    prstGeom = OxmlElement('a:prstGeom')
    prstGeom.set('prst', 'rect')
    spPr.append(prstGeom)
    pic.append(spPr)

    graphicData.append(pic)
    graphic.append(graphicData)
    inline.append(graphic)
    drawing.append(inline)
    r.append(drawing)
    p.append(r)

    return p


def build_caption_para(text):
    """Create a centered italic caption paragraph (OxmlElement)."""
    p = OxmlElement('w:p')
    pPr = OxmlElement('w:pPr')
    jc = OxmlElement('w:jc')
    jc.set(qn('w:val'), 'center')
    pPr.append(jc)
    p.append(pPr)
    p.append(make_run(text, italic=True, size_pt=10))
    return p


def build_table(headers, rows, font_size=8):
    """Build a compact w:tbl element."""
    tbl = OxmlElement('w:tbl')
    ncols = len(headers)

    tblPr = OxmlElement('w:tblPr')
    tblW = OxmlElement('w:tblW')
    tblW.set(qn('w:w'), '5000')
    tblW.set(qn('w:type'), 'pct')
    tblPr.append(tblW)

    tblBorders = OxmlElement('w:tblBorders')
    for bname in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        b = OxmlElement(f'w:{bname}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '2')
        b.set(qn('w:space'), '0')
        b.set(qn('w:color'), '999999')
        tblBorders.append(b)
    tblPr.append(tblBorders)

    tblLook = OxmlElement('w:tblLook')
    tblLook.set(qn('w:firstRow'), '1')
    tblPr.append(tblLook)
    tbl.append(tblPr)

    col_width_pct = str(int(5000 / ncols))

    def make_cell_paragraphs(text, bold=False, color=None, bg=None):
        """Build a <w:tc> with tight spacing."""
        tc = OxmlElement('w:tc')
        tcPr = OxmlElement('w:tcPr')
        tcW = OxmlElement('w:tcW')
        tcW.set(qn('w:type'), 'pct')
        tcW.set(qn('w:w'), col_width_pct)
        tcPr.append(tcW)
        if bg:
            shading = OxmlElement('w:shd')
            shading.set(qn('w:fill'), bg)
            shading.set(qn('w:val'), 'clear')
            tcPr.append(shading)
        # Tight cell margins
        tcMar = OxmlElement('w:tcMar')
        for side in ['top', 'left', 'bottom', 'right']:
            m = OxmlElement(f'w:{side}')
            m.set(qn('w:w'), '30')
            m.set(qn('w:type'), 'dxa')
            tcMar.append(m)
        tcPr.append(tcMar)
        tc.append(tcPr)

        cp = OxmlElement('w:p')
        cpPr = OxmlElement('w:pPr')
        spacing = OxmlElement('w:spacing')
        spacing.set(qn('w:line'), '240')
        spacing.set(qn('w:lineRule'), 'auto')
        spacing.set(qn('w:before'), '0')
        spacing.set(qn('w:after'), '0')
        cpPr.append(spacing)
        if bold:
            jc = OxmlElement('w:jc')
            jc.set(qn('w:val'), 'center')
            cpPr.append(jc)
        cp.append(cpPr)
        cp.append(make_run(str(text), bold=bold, size_pt=font_size, color=color))
        tc.append(cp)
        return tc

    # Header row
    tr_h = OxmlElement('w:tr')
    for h in headers:
        tr_h.append(make_cell_paragraphs(h, bold=True, color='FFFFFF', bg='1A73E8'))
    tbl.append(tr_h)

    # Data rows
    for ri, row in enumerate(rows):
        tr = OxmlElement('w:tr')
        bg = 'F5F5F5' if ri % 2 == 1 else None
        for val in row:
            tr.append(make_cell_paragraphs(val, bg=bg))
        tbl.append(tr)

    return tbl


def insert_before(ref_para, new_element):
    """ref_para can be a Paragraph (has _p) or OxmlElement (already a w:p CT_P)."""
    if hasattr(ref_para, '_p'):
        ref_para._p.addprevious(new_element)
    else:
        ref_para.addprevious(new_element)


def insert_after(ref_para, new_element):
    if hasattr(ref_para, '_p'):
        ref_para._p.addnext(new_element)
    else:
        ref_para.addnext(new_element)


def replace_paragraph_text(para, new_text):
    for r in list(para._p.findall(qn('w:r'))):
        para._p.remove(r)
    r_elem = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), FONT_NAME)
    rFonts.set(qn('w:hAnsi'), FONT_NAME)
    rPr.append(rFonts)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), '24')
    rPr.append(sz)
    r_elem.append(rPr)
    t = OxmlElement('w:t')
    t.text = new_text
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    r_elem.append(t)
    para._p.append(r_elem)


# ── FIGURES ── (processed in REVERSE order)
FIGURES = [
    # (search_text_in_para, image_file, caption_text)
    ("Figure 1: Organigramme", "organigramme_tagip.png",
     "Figure 1 : Organigramme simplifié de TAG-IP avec localisation du poste de stage"),
    ("Figure 2 : Schéma de l", "architecture_reseau.png",
     "Figure 2 : Schéma de l'architecture réseau et cycle de traitement des flux de données"),
    ("Figure 3 : Architecture globale", "architecture_globale.png",
     "Figure 3 : Architecture globale du système (Ash Framework + Phoenix + PostgreSQL)"),
    ("Figure 4 :", "mcd_tag_monitor.png",
     "Figure 4 : Modèle Conceptuel de Données (MCD) du système TAG-Monitor"),
    ("Figure 5 : Modèle Logique de Données (MLD)", "mld_tag_monitor.png",
     "Figure 5 : Modèle Logique de Données (MLD) du système TAG-Monitor"),
    ("Figure 6 : ", "architecture_couches.png",
     "Figure 6 : Schéma d'architecture en couches (Présentation / Métier / Persistance)"),
    ("Figure 7 : Interface du formulaire", "auth_interface.png",
     "Figure 7 : Interface du formulaire d'authentification de TAG-Monitor"),
    ("Figure 8 : Tableau de bord", "dashboard.png",
     "Figure 8 : Tableau de bord de la plateforme de gestion des profils de montage"),
    ("Figure 9 : Interface de configuration initiale", "wizard_etape1.png",
     "Figure 9 : Interface de configuration initiale du profil (Étape 1 - Identification)"),
]


def remove_old_images(doc):
    """Remove only our previously inserted image paragraphs, keeping original captions intact."""
    ns_w = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'

    captions = [
        "Figure 1:", "Figure 2 :", "Figure 3 :",
        "Figure 4 :", "Figure 5 :", "Figure 6 :",
        "Figure 7 :", "Figure 8 :", "Figure 9 :",
    ]

    count = 0
    for p in list(doc.paragraphs):
        text = p.text.strip()
        if any(text.startswith(c) for c in captions):
            p_elem = p._p

            # Remove the image paragraph immediately before the caption
            prev = p_elem.getprevious()
            if prev is not None and prev.tag == f'{{{ns_w}}}p':
                drawing = prev.find(f'.//{{{ns_w}}}drawing')
                if drawing is not None:
                    prev.getparent().remove(prev)
                    count += 1
                else:
                    # Maybe there's a spacer between image and caption
                    prev2 = prev.getprevious()
                    if prev2 is not None and prev2.tag == f'{{{ns_w}}}p':
                        drawing2 = prev2.find(f'.//{{{ns_w}}}drawing')
                        if drawing2 is not None:
                            prev2.getparent().remove(prev2)
                            count += 1

    if count:
        print(f"  🧹 Nettoyé {count} anciens dessins")


def remove_old_tables(doc):
    """Remove only our previously inserted tables and captions, leaving original tables intact."""
    ns_w = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'

    captions = [
        "Tableau 1 : Composants de l'infrastructure matérielle",
        "Tableau 2 : Acteurs du système et droits d'accès",
        "Tableau 3 : Modules fonctionnels de TAG-Monitor",
        "Tableau 4 : Composants HEEx et design system",
        "Tableau 5 : Choix technologiques de TAG-Monitor",
        "Tableau 6 : Tests fonctionnels des LiveViews",
        "Tableau 7 : Variables d'environnement de TAG-Monitor",
    ]

    removed_caps = 0
    removed_tbls = 0
    for p in list(doc.paragraphs):
        if p.text.strip() in captions:
            p_elem = p._p
            parent = p_elem.getparent()

            # Remove the w:tbl that precedes our caption
            prev = p_elem.getprevious()
            if prev is not None and prev.tag == f'{{{ns_w}}}tbl':
                parent.remove(prev)
                removed_tbls += 1

            # Remove the empty spacer paragraph after our caption
            nxt = p_elem.getnext()
            if nxt is not None and nxt.tag == f'{{{ns_w}}}p':
                nxt_texts = nxt.findall(f'.//{{{ns_w}}}t')
                if not nxt_texts or all(
                    t.text is None or t.text.strip() == '' for t in nxt_texts
                ):
                    parent.remove(nxt)

            # Remove the caption paragraph itself
            parent.remove(p_elem)
            removed_caps += 1

    if removed_caps or removed_tbls:
        print(f"  🧹 Nettoyé {removed_caps} légendes + {removed_tbls} tableaux")


def process():
    doc = Document(SRC)
    ns_w = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'

    # ══════════════════════════════
    # 0. Remove only our previously inserted tables (if any)
    # ══════════════════════════════
    remove_old_tables(doc)

    # ══════════════════════════════
    # 1. INSERT / REPLACE IMAGES
    # ══════════════════════════════
    # For each figure: if our image file exists, replace whatever is before
    # the caption; if not, leave the original image in place.
    images_done = 0
    images_skipped = 0
    for search, img_file, caption in reversed(FIGURES):
        ref = None
        for p in doc.paragraphs:
            if search in p.text:
                ref = p
                break
        if ref is None:
            print(f"  ⚠️  Not found: \"{search}\"")
            continue

        img_path = os.path.join(IMG_DIR, img_file)
        if not os.path.exists(img_path):
            print(f"  ⚠️  Image missing: {img_path}")
            images_skipped += 1
            continue

        # Remove any existing image paragraph immediately before the caption
        prev = ref._p.getprevious()
        if prev is not None and prev.tag == f'{{{ns_w}}}p':
            drawing = prev.find(f'.//{{{ns_w}}}drawing')
            if drawing is not None:
                prev.getparent().remove(prev)

        # Build and insert our new image
        try:
            img_p = build_image_para(doc, img_path, width_inches=3.8)
        except Exception as e:
            print(f"  ❌ Error: {img_file}: {e}")
            continue

        insert_before(ref, img_p)
        images_done += 1
        print(f"  ✅ {img_file}")

    # ══════════════════════════════
    # 2. Extract embedded captions + fix formatting
    # ══════════════════════════════
    # These figure captions are embedded inside longer paragraphs.
    # After the image was inserted, we extract the caption text
    # into its own clean paragraph and strip it from the original.

    fixes = [
        # (search_in_para, extract_phrase, clean_caption_text)
        ("RG-03 : Les relations many-to-many",
         "Figure 4 :",
         "Figure 4 : Modèle Conceptuel de Données (MCD) du système TAG-Monitor"),
        ("Les requetes SQL optimisees avec lazy loading",
         "Figure 6",
         "Figure 6 : Schéma d'architecture en couches (Présentation / Métier / Persistance)"),
    ]

    for search, extract_phrase, clean in fixes:
        for p in doc.paragraphs:
            p_lower = p.text.lower()
            if search.lower() in p_lower and extract_phrase.lower() in p_lower:
                # Keep everything before the figure caption as the clean paragraph
                idx = p_lower.find(extract_phrase.lower())
                prefix = p.text[:idx].rstrip(",-;: \t")

                # Create clean caption paragraph
                cap_p = OxmlElement('w:p')
                cap_p.append(make_run(clean, italic=True, size_pt=10))
                p._p.addprevious(cap_p)

                # Replace original paragraph text with just the prefix
                replace_paragraph_text(p, prefix)

                print(f"  ✅ Extracted caption: {clean}")
                break

    # Fix Figure 7/8 combined
    for p in doc.paragraphs:
        if "Figure 7" in p.text and "figure 8" in p.text.lower():
            cap8 = ("Figure 8 : Tableau de bord de la plateforme "
                    "de gestion des profils de montage")
            fig8_p = OxmlElement('w:p')
            fig8_p.append(make_run(cap8, italic=True, size_pt=10))
            insert_after(p, fig8_p)
            replace_paragraph_text(
                p,
                "Figure 7 : Interface du formulaire d'authentification de TAG-Monitor")
            print("  ✅ Figure 7/8 split")
            break

    # Fix "figure 6" lowercase references
    fix6 = 0
    for p in doc.paragraphs:
        for r in p.runs:
            if "figure 6" in r.text.lower() and "Figure 6" not in r.text:
                r.text = r.text.replace("figure 6", "Figure 6")
                fix6 += 1
    if fix6:
        print(f"  ✅ Fixed 'figure 6' → 'Figure 6'")

    # Fix Figure 9 extra text
    for p in doc.paragraphs:
        if p.text.strip().endswith("étape 1 - Identification du profil:"):
            replace_paragraph_text(
                p,
                "Figure 9 : Interface de configuration initiale du profil "
                "(Étape 1 - Identification)")
            print("  ✅ Figure 9 caption fixed")
            break

    # ══════════════════════════════
    # 5. INSERT TABLES (reverse order)
    # ══════════════════════════════
    TABLES = [
        ("6.5.2 Configuration et variables",
         "Tableau 7 : Variables d'environnement de TAG-Monitor",
         ["Variable", "Description", "Valeur par défaut"],
         [["PORT", "Port d'écoute du serveur Phoenix", "4000"],
          ["SECRET_KEY_BASE", "Clé de signature des cookies", "Générée automatiquement"],
          ["POOL_SIZE", "Taille du pool de connexions DB", "10"],
          ["DATABASE_URL", "Chaîne de connexion PostgreSQL", "Variable d'environnement"]]),
        ("tests fonctionnels exécutés",
         "Tableau 6 : Tests fonctionnels des LiveViews",
         ["Test", "Scénario", "Résultat"],
         [["Navigation", "Accès aux pages après authentification", "OK"],
          ["Création profil", "Assistant 4 étapes", "OK"],
          ["Calcul compatibilité", "Moteur 20 critères", "OK"],
          ["Recherche", "Filtres multicritères", "OK"]]),
        ("Le tableau suivant présente les choix technologiques",
         "Tableau 5 : Choix technologiques de TAG-Monitor",
         ["Couche", "Technologie", "Justification"],
         [["Langage", "Elixir", "Concurrence, tolérance aux pannes"],
          ["Framework web", "Phoenix LiveView", "Rendu côté serveur, temps réel"],
          ["Framework métier", "Ash Framework", "Déclaratif, génération automatique"],
          ["Base de données", "PostgreSQL", "Robustesse, géospatial (PostGIS)"],
          ["CSS", "Tailwind CSS", "Productivité, design system"]]),
        ("Composants HEEx et design system",
         "Tableau 4 : Composants HEEx et design system",
         ["Composant", "Rôle", "Fonctionnalités"],
         [["Carte de score", "Affichage résultat compatibilité", "Code couleur, barre progression"],
          ["Carte statistique", "Dashboard temps réel", "Mise à jour PubSub"],
          ["Formulaire", "Saisie profil montage", "Validation progressive"],
          ["Liste", "Catalogue modèles", "Recherche, pagination"]]),
        ("Le système TAG-Monitor est organisé en modules",
         "Tableau 3 : Modules fonctionnels de TAG-Monitor",
         ["Module", "Ressources Ash", "Fonctionnalités"],
         [["ProfilMontage", "ProfilMontage", "CRUD profils, assistant multi-étapes"],
          ["ModeleTraceur", "ModeleTraceur, TypeVehicule", "Catalogue, relations N-N"],
          ["Compatibilité", "Compatibilite", "Moteur de scoring 20 critères"],
          ["Référentiels", "PortType, Feature, Peripheral", "Données de base"],
          ["Dashboard", "Toutes", "Statistiques temps réel via PubSub"]]),
        ("Acteurs du système",
         "Tableau 2 : Acteurs du système et droits d'accès",
         ["Acteur", "Rôle", "Droits"],
         [["Administrateur", "Gestion complète", "CRUD toutes ressources"],
          ["Technicien exploitation", "Préparation et diagnostic", "CRUD profils, consultation"],
          ["Technicien terrain", "Installation sur site", "Consultation fiches profils"]]),
        ("segmenté en trois VLAN",
         "Tableau 1 : Composants de l'infrastructure matérielle",
         ["Composant", "Description", "Rôle"],
         [["Liaisons spécialisées", "Connexions haut débit Telma",
           "Réception trames GPRS/3G"],
          ["VLAN (3)", "Administration, GPS, Développement",
           "Segmentation et sécurité"],
          ["Firewall", "Professionnel IDS/IPS",
           "Sécurité périmétrique"],
          ["Cluster serveurs", "Physiques en HA",
           "Haute disponibilité"],
          ["Stockage RAID", "Réplication disques",
           "Sauvegarde temps réel"]]),
    ]

    tables_done = 0
    for search, caption, headers, rows in reversed(TABLES):
        ref = None
        for p in doc.paragraphs:
            if search in p.text:
                ref = p
                break
        if ref is None:
            print(f"  ⚠️  Table anchor not found: \"{search}\"")
            continue

        # Insert after anchor: table > caption > blank line
        spacer = make_empty_para()
        ref._p.addnext(spacer)

        cap_p = make_empty_para()
        cap_p.append(make_run(caption, bold=True, size_pt=10))
        ref._p.addnext(cap_p)

        tbl_elem = build_table(headers, rows)
        ref._p.addnext(tbl_elem)

        tables_done += 1
        print(f"  ✅ {caption}")

    # ══════════════════════════════
    # SAVE
    # ══════════════════════════════
    doc.save(DST)
    print(f"\n✅ Saved: {DST}")
    print(f"📸 {images_done} images insérées / remplacées")
    if images_skipped:
        print(f"⚠️  {images_skipped} captures manquantes (images originales conservées)")
    print(f"📊 {tables_done} tableaux insérés")


if __name__ == "__main__":
    process()
